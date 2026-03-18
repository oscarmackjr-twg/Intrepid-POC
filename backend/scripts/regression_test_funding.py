"""Funding & CashFlow regression test harness.

Runs only the Final Funding (SG + CIBC) and CashFlow phases against each
buy-date folder in TestData and diffs outputs against golden expected files.
Assumes files_required/ already contains the split _sg/_cibc exhibit files
(i.e. tagging has already been run and its outputs are part of the golden data).

Phases run per test case:
  1. Funding SG   — scripts/final_funding_sg.py
  2. Funding CIBC — scripts/final_funding_cibc.py
  3. CashFlow     — cashflow/compute/run_purchase_package.py (SG and CIBC)

Each test case runs against a temp copy of files_required/ so the golden
TestData directory is never modified.  Generated outputs land in
temp_dir/output/ and temp_dir/output_share/, which are then diffed against the
golden test_case_dir/output/ and test_case_dir/output_share/.

Usage (from repo root):
    python backend/scripts/regression_test_funding.py
    python backend/scripts/regression_test_funding.py --test-data C:\\Users\\omack\\Downloads\\TestData
    python backend/scripts/regression_test_funding.py --pdate 2026-02-24 --tday 2026-02-19
    python backend/scripts/regression_test_funding.py --no-cleanup
    python backend/scripts/regression_test_funding.py --report report.xlsx

TestData folder structure:
    <test-data-dir>/
        {buy_date_folder}/
            files_required/     # must include _sg and _cibc split exhibit files
            output/
                <expected output files>
            output_share/
                <expected output_share files>
        dates.json              # optional per-folder date config

Exit code: 0 if all cases PASS, 1 if any case FAILS.
"""
import argparse
import filecmp
import fnmatch
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path


def _add_us_business_days(tday_str: str, n: int) -> str:
    """Add n US business days to tday (YYYY-MM-DD), skipping weekends and US holidays."""
    try:
        import holidays as _hl
        d = datetime.strptime(tday_str, "%Y-%m-%d").date()
        us_hols = _hl.US(years=range(d.year, d.year + 2))
        count = 0
        while count < n:
            d += timedelta(days=1)
            if d.weekday() < 5 and d not in us_hols:
                count += 1
        return d.strftime("%Y-%m-%d")
    except ImportError:
        d = datetime.strptime(tday_str, "%Y-%m-%d").date()
        count = 0
        while count < n:
            d += timedelta(days=1)
            if d.weekday() < 5:
                count += 1
        return d.strftime("%Y-%m-%d")


# ---------------------------------------------------------------------------
# Path helpers
# ---------------------------------------------------------------------------

def _repo_root() -> Path:
    return Path(__file__).resolve().parent.parent.parent


def _default_backend_dir() -> Path:
    return _repo_root() / "backend"


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def discover_test_cases(test_data_dir: Path) -> list[Path]:
    """Return sorted list of valid test-case directories under test_data_dir."""
    if not test_data_dir.exists():
        print(f"[WARN] Test data directory not found: {test_data_dir}", file=sys.stderr)
        return []

    cases = []
    for subdir in sorted(test_data_dir.iterdir()):
        if not subdir.is_dir():
            continue
        files_required = subdir / "files_required"
        has_files = files_required.is_dir() and any(f.is_file() for f in files_required.iterdir())
        has_expected = (subdir / "output").is_dir() or (subdir / "output_share").is_dir()
        if not has_files:
            print(f"[WARN] Skipping {subdir.name}: no input files found in files_required/")
            continue
        if not has_expected:
            print(f"[WARN] Skipping {subdir.name}: no output/ or output_share/ directory")
            continue

        # Check for split exhibit files (_sg / _cibc) — required for this harness
        fr = subdir / "files_required"
        sg_files  = list(fr.glob("*_sg.xlsx"))
        cibc_files = list(fr.glob("*_cibc.xlsx"))
        if not sg_files or not cibc_files:
            print(
                f"[WARN] Skipping {subdir.name}: no _sg.xlsx or _cibc.xlsx exhibit files "
                "found in files_required/ (run tagging first)"
            )
            continue

        cases.append(subdir)

    return cases


# ---------------------------------------------------------------------------
# Date helpers
# ---------------------------------------------------------------------------

def _load_dates_config(test_data_dir: Path) -> dict:
    config_path = test_data_dir / "dates.json"
    if not config_path.exists():
        return {}
    try:
        with open(config_path) as f:
            return json.load(f)
    except Exception as exc:
        print(f"[WARN] Could not read dates.json: {exc}")
        return {}


def _derive_dates(
    folder_name: str,
    cli_pdate: str | None,
    cli_tday: str | None,
    dates_config: dict | None = None,
) -> tuple[str, str]:
    """Return (pdate, tday) for a test case.

    tday is resolved from: CLI arg → dates.json → today.
    pdate is always tday + 3 US business days unless overridden via --pdate.
    """
    today_str = date.today().isoformat()

    if cli_tday:
        tday = cli_tday
    elif dates_config and folder_name in dates_config:
        tday = dates_config[folder_name].get("tday", today_str)
    else:
        tday = today_str

    pdate = cli_pdate if cli_pdate else _add_us_business_days(tday, 3)

    return pdate, tday


def _derive_date_vars(tday: str) -> dict:
    tday_dt = datetime.strptime(tday, "%Y-%m-%d")
    yesterday_dt = tday_dt - timedelta(days=1)
    curr_date = tday_dt.strftime("%m-%d-%Y")
    yesterday_str = yesterday_dt.strftime("%m-%d-%Y")
    first_of_month = tday_dt.replace(day=1)
    last_of_prev_month = first_of_month - timedelta(days=1)
    last_end = f"{last_of_prev_month.year}_{last_of_prev_month.month:03}_{last_of_prev_month.day:02}"
    fd = first_of_month.strftime("%Y-%m-%d")
    return {
        "curr_date":  curr_date,
        "yesterday":  yesterday_str,
        "last_end":   last_end,
        "fd":         fd,
    }


# ---------------------------------------------------------------------------
# Phase classification helpers
# ---------------------------------------------------------------------------

PHASE_COLS = [
    "Final Funding\nSG",
    "Final Funding\nCIBC",
    "CashFlow",
]
PHASE_KEYS = ["funding_sg", "funding_cibc", "cashflow"]


def _classify_output_file(filename: str) -> str:
    name = Path(filename).name
    nl = name.lower()

    for pat in ["sfc_cashflows_*", "twg_cashflows_*", "loans_data_*",
                "cashflows_*", "cashflow profile*"]:
        if fnmatch.fnmatch(nl, pat):
            return "cashflow"

    for pat in ["everyloan_sg*", "borrowing_file_sg*", "concentration_final_sg*",
                "flagged_loans_*_sg*", "notes_flagged_loans_*_sg*"]:
        if fnmatch.fnmatch(nl, pat):
            return "funding_sg"

    for pat in ["everyloan_cibc*", "borrowing_file_cibc*", "concentration_final_cibc*",
                "flagged_loans_*_cibc*", "notes_flagged_loans_*_cibc*"]:
        if fnmatch.fnmatch(nl, pat):
            return "funding_cibc"

    for pat in ["comap_not_passed*", "purchase_price_mismatch*",
                "flagged_loans_*", "notes_flagged_loans_*"]:
        if fnmatch.fnmatch(nl, pat):
            return "funding_sg"

    return "funding_sg"  # default


# ---------------------------------------------------------------------------
# Diff logic
# ---------------------------------------------------------------------------

def _collect_files(directory: Path) -> dict[str, Path]:
    result = {}
    if not directory.exists():
        return result
    for f in directory.rglob("*"):
        if f.is_file():
            rel = str(f.relative_to(directory))
            result[rel] = f
    return result


def _diff_directories(
    generated_dir: Path,
    expected_dir: Path,
    label: str,
) -> tuple[list[str], list[str], list[str]]:
    diffs: list[str] = []
    missing: list[str] = []
    extras: list[str] = []

    if not expected_dir.exists():
        return diffs, missing, extras

    expected_files  = _collect_files(expected_dir)
    generated_files = _collect_files(generated_dir) if generated_dir.exists() else {}

    for rel, exp_path in expected_files.items():
        display = f"{label}/{rel}"
        if rel not in generated_files:
            missing.append(display)
        else:
            try:
                same = filecmp.cmp(str(generated_files[rel]), str(exp_path), shallow=False)
            except OSError as exc:
                diffs.append(f"{display} (compare error: {exc})")
                continue
            if not same:
                diffs.append(display)

    for rel in generated_files:
        if rel not in expected_files:
            extras.append(f"{label}/{rel}")

    return diffs, missing, extras


def _compare_tabular_files(actual_path: Path, expected_path: Path) -> list[dict]:
    import pandas as pd

    KEY_COLS = ["SELLER Loan #", "Account Number", "Loan #", "loan_num", "dates"]
    MAX_DIFFS = 200

    def _load(p: Path):
        try:
            if p.suffix.lower() in (".xlsx", ".xls"):
                return pd.read_excel(p, dtype=str)
            elif p.suffix.lower() == ".csv":
                return pd.read_csv(p, dtype=str)
        except Exception:
            pass
        return None

    actual   = _load(actual_path)
    expected = _load(expected_path)

    if actual is None or expected is None:
        return [{"loan_num": "—", "column": "(binary)",
                 "expected": "—", "actual": "—",
                 "note": "cannot compare non-tabular file"}]

    # everyloan files have an auto-generated row-number first column; skip it
    import fnmatch as _fnmatch
    if _fnmatch.fnmatch(actual_path.name.lower(), "everyloan_*"):
        if len(actual.columns) > 1:
            actual   = actual.iloc[:, 1:]
        if len(expected.columns) > 1:
            expected = expected.iloc[:, 1:]

    # cashflow CSVs have (loan_number, dates) as compound key — sort both so
    # positional comparison is loan-order-independent; force positional (skip key lookup)
    _force_positional = False
    if actual_path.suffix.lower() == ".csv" and "loan_number" in actual.columns and "dates" in actual.columns:
        sort_cols = [c for c in ["loan_number", "dates"] if c in actual.columns and c in expected.columns]
        if sort_cols:
            actual   = actual.sort_values(sort_cols).reset_index(drop=True)
            expected = expected.sort_values(sort_cols).reset_index(drop=True)
            _force_positional = True

    actual   = actual.fillna("").astype(str)
    expected = expected.fillna("").astype(str)

    diffs: list[dict] = []
    key_col = None if _force_positional else next((k for k in KEY_COLS if k in actual.columns and k in expected.columns), None)

    if key_col is None:
        if actual.shape != expected.shape:
            diffs.append({
                "loan_num": "—", "column": "(shape)",
                "expected": f"{expected.shape[0]} rows × {expected.shape[1]} cols",
                "actual":   f"{actual.shape[0]} rows × {actual.shape[1]} cols",
                "note": "row/column count differs",
            })
            return diffs
        for i in range(len(expected)):
            for col in expected.columns:
                if col not in actual.columns:
                    continue
                ev = str(expected.iloc[i][col]).strip()
                av = str(actual.iloc[i][col]).strip()
                if ev != av:
                    diffs.append({"loan_num": f"row {i+1}", "column": col,
                                  "expected": ev, "actual": av, "note": ""})
                    if len(diffs) >= MAX_DIFFS:
                        return diffs
        return diffs

    try:
        act_idx = actual.set_index(key_col)
        exp_idx = expected.set_index(key_col)
    except Exception:
        return [{"loan_num": "—", "column": "(index)",
                 "expected": "—", "actual": "—",
                 "note": "could not index on key column"}]

    for loan in exp_idx.index.difference(act_idx.index):
        diffs.append({"loan_num": str(loan), "column": "(row)",
                      "expected": "present", "actual": "missing",
                      "note": "loan absent from output"})
        if len(diffs) >= MAX_DIFFS:
            return diffs

    for loan in act_idx.index.difference(exp_idx.index):
        diffs.append({"loan_num": str(loan), "column": "(row)",
                      "expected": "absent", "actual": "present",
                      "note": "unexpected loan in output"})
        if len(diffs) >= MAX_DIFFS:
            return diffs

    common_loans = exp_idx.index.intersection(act_idx.index)
    common_cols  = [c for c in exp_idx.columns if c in act_idx.columns]
    for loan in common_loans:
        exp_row = exp_idx.loc[loan]
        act_row = act_idx.loc[loan]
        for col in common_cols:
            ev = str(exp_row[col]).strip()
            av = str(act_row[col]).strip()
            if ev != av:
                diffs.append({"loan_num": str(loan), "column": col,
                              "expected": ev, "actual": av, "note": ""})
                if len(diffs) >= MAX_DIFFS:
                    return diffs

    return diffs


# ---------------------------------------------------------------------------
# Phase runner
# ---------------------------------------------------------------------------

def _run_phase(
    phase_name: str,
    cmd: list[str],
    cwd: str,
    env: dict | None = None,
    timeout: int = 300,
) -> tuple[bool, str | None]:
    run_env = os.environ.copy()
    if env:
        run_env.update(env)

    try:
        proc = subprocess.run(
            cmd,
            cwd=cwd,
            capture_output=True,
            text=True,
            timeout=timeout,
            env=run_env,
        )
        if proc.stdout:
            for line in proc.stdout.splitlines():
                print(f"  [{phase_name}] {line}")
        if proc.stderr:
            for line in proc.stderr.splitlines():
                print(f"  [{phase_name} STDERR] {line}")

        if proc.returncode != 0:
            return False, f"{phase_name} exited with code {proc.returncode}"
        return True, None

    except subprocess.TimeoutExpired:
        return False, f"{phase_name} timed out after {timeout}s"
    except Exception as exc:
        return False, f"{phase_name} exception: {exc}"


# ---------------------------------------------------------------------------
# Per-test-case runner
# ---------------------------------------------------------------------------

def run_test_case(
    test_case_dir: Path,
    backend_dir: Path,
    cli_pdate: str | None,
    cli_tday: str | None,
    no_cleanup: bool,
    dates_config: dict | None = None,
    update_golden: bool = False,
) -> dict:
    pdate, tday = _derive_dates(test_case_dir.name, cli_pdate, cli_tday, dates_config)
    date_vars = _derive_date_vars(tday)
    buy_num = test_case_dir.name.replace("_buy", "")

    result = {
        "name":          test_case_dir.name,
        "test_case_dir": test_case_dir,
        "pdate":         pdate,
        "tday":          tday,
        "buy_num":       buy_num,
        "date_vars":     date_vars,
        "status":        "FAILED",
        "phases_ok":     [],
        "phases_failed": [],
        "diffs":         [],
        "missing":       [],
        "extra":         [],
        "exceptions":    [],
        "error":         None,
        "work_dir":      None,
    }

    print(f"\n[BUY DATE: {test_case_dir.name}]")
    print(f"  pdate={pdate}  tday={tday}  curr_date={date_vars['curr_date']}  buy_num={buy_num}")

    tmp_parent = Path(tempfile.mkdtemp(prefix="regression_funding_"))
    work_dir = tmp_parent / "work"
    work_dir.mkdir()
    result["work_dir"] = work_dir

    try:
        shutil.copytree(str(test_case_dir / "files_required"), str(work_dir / "files_required"))
        (work_dir / "output").mkdir()
        (work_dir / "output_share").mkdir()
        print(f"  Work dir: {work_dir}")

        funding_env = {
            "FOLDER":     str(work_dir),
            "PDATE":      pdate,
            "CURR_DATE":  date_vars["curr_date"],
            "YESTERDAY":  date_vars["yesterday"],
            "LAST_END":   date_vars["last_end"],
            "FD":         date_vars["fd"],
            "BUY_NUM":    buy_num,
        }

        # Phase 1: Funding SG
        print("  [Phase 1] Funding SG...")
        ok, err = _run_phase(
            "FundingSG",
            [sys.executable, "scripts/final_funding_sg.py"],
            cwd=str(backend_dir),
            env=funding_env,
            timeout=300,
        )
        if ok:
            result["phases_ok"].append("Phase1:FundingSG")
            print("  [Phase 1] OK")
        else:
            result["phases_failed"].append(f"Phase1:FundingSG — {err}")
            result["error"] = err
            print(f"  [Phase 1] FAILED: {err}")
            return result

        # Phase 2: Funding CIBC
        print("  [Phase 2] Funding CIBC...")
        ok, err = _run_phase(
            "FundingCIBC",
            [sys.executable, "scripts/final_funding_cibc.py"],
            cwd=str(backend_dir),
            env=funding_env,
            timeout=300,
        )
        if ok:
            result["phases_ok"].append("Phase2:FundingCIBC")
            print("  [Phase 2] OK")
        else:
            result["phases_failed"].append(f"Phase2:FundingCIBC — {err}")
            result["error"] = err
            print(f"  [Phase 2] FAILED: {err}")
            return result

        # Phase 3: CashFlow (SG and CIBC)
        files_req  = work_dir / "files_required"
        output_dir = work_dir / "output"
        curr_date  = date_vars["curr_date"]

        for buyer in ("sg", "cibc"):
            sfy_file   = files_req / f"FX3_{curr_date}_ExhibitAtoFormofSaleNotice_{buyer}.xlsx"
            prime_file = files_req / f"{curr_date} Exhibit A To Form Of Sale Notice_{buyer}.xlsx"

            if not sfy_file.exists() or not prime_file.exists():
                print(f"  [Phase 3 {buyer.upper()}] Skipped — exhibit files not found")
                continue

            print(f"  [Phase 3 {buyer.upper()}] CashFlow...")
            ok, err = _run_phase(
                f"CashFlow-{buyer.upper()}",
                [
                    sys.executable, "-m", "cashflow.compute.run_purchase_package",
                    "--prime-file",    str(prime_file),
                    "--sfy-file",      str(sfy_file),
                    "--master-sheet",  str(files_req / "MASTER_SHEET.xlsx"),
                    "--notes-sheet",   str(files_req / "MASTER_SHEET - Notes.xlsx"),
                    "--purchase-date", pdate,
                    "--output-dir",    str(output_dir),
                    "--buy-num",       buy_num,
                    "--buyer",         buyer,
                ],
                cwd=str(backend_dir),
                timeout=300,
            )
            if ok:
                result["phases_ok"].append(f"Phase3:CashFlow{buyer.upper()}")
                print(f"  [Phase 3 {buyer.upper()}] OK")
            else:
                result["phases_failed"].append(f"Phase3:CashFlow{buyer.upper()} — {err}")
                print(f"  [Phase 3 {buyer.upper()}] FAILED: {err}")

        # --update-golden: overwrite golden expected files with actual outputs
        if update_golden:
            updated_files = []
            for subdir in ("output", "output_share"):
                actual_dir  = work_dir / subdir
                golden_dir  = test_case_dir / subdir
                golden_dir.mkdir(exist_ok=True)
                for actual_file in actual_dir.rglob("*"):
                    if not actual_file.is_file():
                        continue
                    rel = actual_file.relative_to(actual_dir)
                    dest = golden_dir / rel
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(str(actual_file), str(dest))
                    updated_files.append(f"{subdir}/{rel}")
            print(f"  [update-golden] Wrote {len(updated_files)} files to {test_case_dir}")
            for f in updated_files:
                print(f"    {f}")
            result["status"] = "PASS"
            return result

        # Diff outputs vs golden
        expected_outputs = test_case_dir / "output"
        if expected_outputs.exists():
            d, m, e = _diff_directories(output_dir, expected_outputs, "output")
            result["diffs"].extend(d)
            result["missing"].extend(m)
            result["extra"].extend(e)

        expected_share = test_case_dir / "output_share"
        if expected_share.exists():
            d, m, e = _diff_directories(work_dir / "output_share", expected_share, "output_share")
            result["diffs"].extend(d)
            result["missing"].extend(m)
            result["extra"].extend(e)

        # Collect cell-level exceptions for differing files
        for rel_path in result["diffs"]:
            parts = Path(rel_path).parts
            if len(parts) < 2:
                continue
            subdir, file_rel = parts[0], str(Path(*parts[1:]))
            if subdir == "output":
                actual_file   = output_dir / file_rel
                expected_file = test_case_dir / "output" / file_rel
            elif subdir == "output_share":
                actual_file   = work_dir / "output_share" / file_rel
                expected_file = test_case_dir / "output_share" / file_rel
            else:
                continue
            file_diffs = _compare_tabular_files(actual_file, expected_file)
            for d in file_diffs:
                d["file"] = rel_path
                result["exceptions"].append(d)

        total_issues = len(result["diffs"]) + len(result["missing"]) + len(result["extra"])
        result["status"] = "PASS" if total_issues == 0 else "FAILED"

        if total_issues == 0:
            print("  Diff result: PASS (0 diffs)")
        else:
            print(f"  Diff result: FAIL ({total_issues} differences)")
            for item in result["diffs"]:
                print(f"    DIFF: {item}")
            for item in result["missing"]:
                print(f"    MISSING: {item}")
            for item in result["extra"]:
                print(f"    EXTRA: {item}")

    except Exception as exc:
        result["error"] = str(exc)
        print(f"  FAILED (exception: {exc})")
    finally:
        if not no_cleanup and result.get("work_dir") is not None:
            try:
                shutil.rmtree(str(tmp_parent))
            except OSError as exc:
                print(f"  [WARN] Could not clean up {tmp_parent}: {exc}")

    return result


# ---------------------------------------------------------------------------
# Excel report generation
# ---------------------------------------------------------------------------

def _file_status(rel_path: str, result: dict) -> str:
    if rel_path in result["diffs"]:
        return "DIFFER"
    if rel_path in result["missing"]:
        return "MISSING"
    if rel_path in result["extra"]:
        return "EXTRA"
    return "MATCH"


def _all_expected_files(result: dict) -> list[str]:
    files = []
    test_case_dir = result.get("test_case_dir")
    if test_case_dir is None:
        return files
    for subdir_label in ("output", "output_share"):
        golden_dir = Path(test_case_dir) / subdir_label
        if golden_dir.exists():
            for f in golden_dir.rglob("*"):
                if f.is_file():
                    rel = f"{subdir_label}/{f.relative_to(golden_dir)}"
                    files.append(rel)
    return files


def write_excel_report(results: list, report_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    FILL_GREEN  = PatternFill("solid", fgColor="C6EFCE")
    FILL_RED    = PatternFill("solid", fgColor="FFC7CE")
    FILL_ORANGE = PatternFill("solid", fgColor="FFCC99")
    FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
    FILL_BLUE   = PatternFill("solid", fgColor="BDD7EE")
    FILL_DKBLUE = PatternFill("solid", fgColor="4472C4")
    FILL_GRAY   = PatternFill("solid", fgColor="EEEEEE")
    FILL_NONE   = PatternFill("none")

    FONT_BOLD  = Font(bold=True)
    FONT_WHITE = Font(bold=True, color="FFFFFF")
    ALIGN_CTR  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ALIGN_LEFT = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    thin   = Side(style="thin", color="AAAAAA")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _cell(ws, row, col, value=None, fill=None, font=None, align=None, border=None):
        c = ws.cell(row=row, column=col, value=value)
        if fill:   c.fill      = fill
        if font:   c.font      = font
        if align:  c.alignment = align
        if border: c.border    = border
        return c

    def _status_fill(status: str) -> PatternFill:
        return {
            "MATCH":   FILL_GREEN,
            "DIFFER":  FILL_ORANGE,
            "MISSING": FILL_RED,
            "EXTRA":   FILL_YELLOW,
            "PASS":    FILL_GREEN,
            "FAILED":  FILL_RED,
        }.get(status, FILL_NONE)

    wb = Workbook()
    wb.remove(wb.active)

    # -----------------------------------------------------------------------
    # Summary sheet
    # -----------------------------------------------------------------------
    sum_ws = wb.create_sheet("Summary")
    sum_ws.column_dimensions["A"].width = 22
    for col_letter in ["B", "C", "D", "E"]:
        sum_ws.column_dimensions[col_letter].width = 18

    title_cell = sum_ws.cell(row=1, column=1, value="Funding & CashFlow Test Summary")
    title_cell.font = Font(bold=True, size=14)
    sum_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(results))

    _cell(sum_ws, 2, 1, "Metric", fill=FILL_BLUE, font=FONT_BOLD, align=ALIGN_CTR, border=BORDER)
    for col_idx, r in enumerate(results, start=2):
        _cell(sum_ws, 2, col_idx, r["name"], fill=FILL_BLUE, font=FONT_BOLD,
              align=ALIGN_CTR, border=BORDER)

    rows_data = [
        ("pdate",          lambda r: r["pdate"]),
        ("tday",           lambda r: r["tday"]),
        ("Test Status",    lambda r: r["status"]),
        ("Phases OK",      lambda r: len(r["phases_ok"])),
        ("Phases Failed",  lambda r: len(r["phases_failed"])),
        ("Output Match",   lambda r: sum(1 for f in _all_expected_files(r) if _file_status(f, r) == "MATCH")),
        ("Output Differ",  lambda r: len(r["diffs"])),
        ("Output Missing", lambda r: len(r["missing"])),
        ("Output Extra",   lambda r: len(r["extra"])),
        ("Total Issues",   lambda r: len(r["diffs"]) + len(r["missing"]) + len(r["extra"])),
    ]

    for row_offset, (label, fn) in enumerate(rows_data, start=3):
        _cell(sum_ws, row_offset, 1, label, fill=FILL_GRAY, font=FONT_BOLD,
              align=ALIGN_LEFT, border=BORDER)
        for col_idx, r in enumerate(results, start=2):
            val = fn(r)
            fill = FILL_NONE
            if label == "Test Status":
                fill = _status_fill(str(val))
            elif label == "Total Issues":
                fill = FILL_GREEN if val == 0 else FILL_RED
            _cell(sum_ws, row_offset, col_idx, val, fill=fill, align=ALIGN_CTR, border=BORDER)

    # Phase breakdown
    phase_row_start = len(rows_data) + 4
    sum_ws.cell(row=phase_row_start, column=1, value="Phases").font = Font(bold=True, italic=True)

    all_phases: list[str] = []
    for r in results:
        for p in r["phases_ok"] + [p.split(" — ")[0] for p in r["phases_failed"]]:
            if p not in all_phases:
                all_phases.append(p)

    for row_offset, phase in enumerate(all_phases, start=phase_row_start + 1):
        _cell(sum_ws, row_offset, 1, phase, fill=FILL_GRAY, align=ALIGN_LEFT, border=BORDER)
        fail_names = [p.split(" — ")[0] for p in r["phases_failed"]]
        for col_idx, r in enumerate(results, start=2):
            fail_names = [p.split(" — ")[0] for p in r["phases_failed"]]
            if phase in r["phases_ok"]:
                _cell(sum_ws, row_offset, col_idx, "OK",     fill=FILL_GREEN, align=ALIGN_CTR, border=BORDER)
            elif phase in fail_names:
                _cell(sum_ws, row_offset, col_idx, "FAILED", fill=FILL_RED,   align=ALIGN_CTR, border=BORDER)
            else:
                _cell(sum_ws, row_offset, col_idx, "—",      fill=FILL_NONE,  align=ALIGN_CTR, border=BORDER)

    # -----------------------------------------------------------------------
    # Per-test-case sheets
    # -----------------------------------------------------------------------
    for result in results:
        test_case_dir = result["test_case_dir"]
        sheet_name    = result["name"][:31]
        ws = wb.create_sheet(sheet_name)

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 48
        for i in range(len(PHASE_COLS)):
            ws.column_dimensions[get_column_letter(3 + i)].width = 18

        # Row 1: column headers
        _cell(ws, 1, 1, "Item",      fill=FILL_BLUE, font=FONT_BOLD, align=ALIGN_CTR, border=BORDER)
        _cell(ws, 1, 2, "File Name", fill=FILL_BLUE, font=FONT_BOLD, align=ALIGN_LEFT, border=BORDER)
        for i, phase_label in enumerate(PHASE_COLS):
            _cell(ws, 1, 3 + i, phase_label,
                  fill=FILL_BLUE, font=FONT_BOLD, align=ALIGN_CTR, border=BORDER)
        ws.row_dimensions[1].height = 36

        # Row 2: test metadata
        meta = (
            f"pdate={result['pdate']}  "
            f"tday={result['tday']}  "
            f"curr_date={result['date_vars']['curr_date']}  "
            f"buy_num={result['buy_num']}"
        )
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=2 + len(PHASE_COLS))
        _cell(ws, 2, 1, "Dates", fill=FILL_GRAY, font=FONT_BOLD, align=ALIGN_CTR)
        meta_cell = ws.cell(row=2, column=2, value=meta)
        meta_cell.font      = Font(italic=True)
        meta_cell.alignment = ALIGN_LEFT

        current_row = 3

        # INPUT FILES section
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=2 + len(PHASE_COLS))
        sec = ws.cell(row=current_row, column=1, value="INPUT FILES (files_required/)")
        sec.fill = FILL_DKBLUE; sec.font = FONT_WHITE; sec.alignment = ALIGN_CTR
        current_row += 1

        files_required_dir = test_case_dir / "files_required"
        input_files = sorted(f.name for f in files_required_dir.iterdir() if f.is_file())

        for idx, fname in enumerate(input_files, start=1):
            # Mark split exhibit files under the phase that consumes them
            nl = fname.lower()
            if ("_sg" in nl or "_cibc" in nl) and ("exhibit" in nl or "fx3_" in nl):
                phase_key = "funding_sg" if "_sg" in nl else "funding_cibc"
            else:
                phase_key = ""
            phase_idx = PHASE_KEYS.index(phase_key) if phase_key in PHASE_KEYS else -1

            _cell(ws, current_row, 1, f"Input {idx}", fill=FILL_GRAY, align=ALIGN_CTR, border=BORDER)
            _cell(ws, current_row, 2, fname,           fill=FILL_GRAY, align=ALIGN_LEFT, border=BORDER)
            for i in range(len(PHASE_COLS)):
                if i == phase_idx:
                    _cell(ws, current_row, 3 + i, "PRESENT",
                          fill=FILL_GREEN, align=ALIGN_CTR, border=BORDER)
                else:
                    _cell(ws, current_row, 3 + i, "",
                          fill=FILL_NONE, align=ALIGN_CTR, border=BORDER)
            current_row += 1

        # EXPECTED OUTPUTS section
        current_row += 1
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=2 + len(PHASE_COLS))
        sec = ws.cell(row=current_row, column=1,
                      value="EXPECTED OUTPUT FILES (output/ and output_share/)")
        sec.fill = FILL_DKBLUE; sec.font = FONT_WHITE; sec.alignment = ALIGN_CTR
        current_row += 1

        expected_output_files = []
        for subdir_label in ("output", "output_share"):
            golden_dir = test_case_dir / subdir_label
            if golden_dir.exists():
                for f in sorted(golden_dir.rglob("*")):
                    if f.is_file():
                        rel = f"{subdir_label}/{f.relative_to(golden_dir)}"
                        expected_output_files.append(rel)

        for idx, rel_path in enumerate(expected_output_files, start=1):
            fname       = Path(rel_path).name
            status      = _file_status(rel_path, result)
            phase_key   = _classify_output_file(fname)
            phase_idx   = PHASE_KEYS.index(phase_key) if phase_key in PHASE_KEYS else -1
            status_fill = _status_fill(status)
            row_fill    = FILL_GRAY if status == "MATCH" else FILL_NONE

            _cell(ws, current_row, 1, f"Output {idx}", fill=row_fill, align=ALIGN_CTR, border=BORDER)
            _cell(ws, current_row, 2, rel_path,         fill=row_fill, align=ALIGN_LEFT, border=BORDER)
            for i in range(len(PHASE_COLS)):
                if i == phase_idx:
                    _cell(ws, current_row, 3 + i, status,
                          fill=status_fill, align=ALIGN_CTR, border=BORDER)
                else:
                    _cell(ws, current_row, 3 + i, "",
                          fill=FILL_NONE, align=ALIGN_CTR, border=BORDER)
            current_row += 1

        # EXTRA FILES section
        extra_files = result["extra"]
        if extra_files:
            current_row += 1
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row, end_column=2 + len(PHASE_COLS))
            sec = ws.cell(row=current_row, column=1,
                          value="EXTRA FILES (generated but not in golden)")
            sec.fill = FILL_DKBLUE; sec.font = FONT_WHITE; sec.alignment = ALIGN_CTR
            current_row += 1

            for idx, rel_path in enumerate(extra_files, start=1):
                fname     = Path(rel_path).name
                phase_key = _classify_output_file(fname)
                phase_idx = PHASE_KEYS.index(phase_key) if phase_key in PHASE_KEYS else -1

                _cell(ws, current_row, 1, f"Extra {idx}", fill=FILL_YELLOW,
                      align=ALIGN_CTR, border=BORDER)
                _cell(ws, current_row, 2, rel_path,        fill=FILL_YELLOW,
                      align=ALIGN_LEFT, border=BORDER)
                for i in range(len(PHASE_COLS)):
                    if i == phase_idx:
                        _cell(ws, current_row, 3 + i, "EXTRA",
                              fill=FILL_YELLOW, align=ALIGN_CTR, border=BORDER)
                    else:
                        _cell(ws, current_row, 3 + i, "",
                              fill=FILL_NONE, align=ALIGN_CTR, border=BORDER)
                current_row += 1

        # PHASE EXECUTION STATUS section
        current_row += 1
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=2 + len(PHASE_COLS))
        sec = ws.cell(row=current_row, column=1, value="PHASE EXECUTION STATUS")
        sec.fill = FILL_DKBLUE; sec.font = FONT_WHITE; sec.alignment = ALIGN_CTR
        current_row += 1

        phase_ok_set = set(result["phases_ok"])
        fail_map     = {p.split(" — ")[0]: p.split(" — ", 1)[1] if " — " in p else ""
                        for p in result["phases_failed"]}

        phase_display = [
            ("Phase 1: Funding SG",   "Phase1:FundingSG"),
            ("Phase 2: Funding CIBC", "Phase2:FundingCIBC"),
            ("Phase 3: CashFlow SG",  "Phase3:CashFlowSG"),
            ("Phase 3: CashFlow CIBC","Phase3:CashFlowCIBC"),
        ]
        for phase_label, phase_key in phase_display:
            if phase_key in phase_ok_set:
                status_val, s_fill = "OK", FILL_GREEN
            elif phase_key in fail_map:
                status_val, s_fill = f"FAILED: {fail_map[phase_key]}", FILL_RED
            else:
                status_val, s_fill = "—", FILL_NONE

            _cell(ws, current_row, 1, phase_label, fill=FILL_GRAY,
                  font=FONT_BOLD, align=ALIGN_LEFT, border=BORDER)
            ws.merge_cells(start_row=current_row, start_column=2,
                           end_row=current_row, end_column=2 + len(PHASE_COLS))
            _cell(ws, current_row, 2, status_val, fill=s_fill, align=ALIGN_LEFT, border=BORDER)
            current_row += 1

        # Overall summary row
        current_row += 1
        total_issues   = len(result["diffs"]) + len(result["missing"]) + len(result["extra"])
        overall_status = "PASS" if result["status"] == "PASS" else f"FAILED ({total_issues} issues)"
        _cell(ws, current_row, 1, "Overall", fill=FILL_GRAY, font=FONT_BOLD,
              align=ALIGN_CTR, border=BORDER)
        ws.merge_cells(start_row=current_row, start_column=2,
                       end_row=current_row, end_column=2 + len(PHASE_COLS))
        _cell(ws, current_row, 2, overall_status,
              fill=_status_fill(result["status"]), font=FONT_BOLD, align=ALIGN_CTR, border=BORDER)

        # EXCEPTIONS section
        exceptions   = result.get("exceptions", [])
        missing_rows = [{"file": f, "loan_num": "—", "column": "(file)",
                         "expected": "present", "actual": "missing",
                         "note": "file not generated"}
                        for f in result.get("missing", [])]
        extra_rows   = [{"file": f, "loan_num": "—", "column": "(file)",
                         "expected": "absent", "actual": "present",
                         "note": "unexpected file generated"}
                        for f in result.get("extra", [])]
        all_exceptions = exceptions + missing_rows + extra_rows

        current_row += 2
        n_cols = 2 + len(PHASE_COLS)
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=n_cols)
        sec = ws.cell(row=current_row, column=1,
                      value=f"EXCEPTIONS / DIFFERENCES  ({len(all_exceptions)} items)")
        sec.fill = FILL_DKBLUE; sec.font = FONT_WHITE; sec.alignment = ALIGN_CTR
        current_row += 1

        exc_headers = ["#", "File", "Loan #", "Column", "Expected Value", "Actual Value", "Notes"]
        for i, h in enumerate(exc_headers, start=1):
            _cell(ws, current_row, i, h, fill=FILL_BLUE, font=FONT_BOLD,
                  align=ALIGN_CTR, border=BORDER)
        current_row += 1

        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 28
        ws.column_dimensions["F"].width = 28
        ws.column_dimensions["G"].width = 34

        if not all_exceptions:
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row, end_column=n_cols)
            _cell(ws, current_row, 1, "No differences found — all outputs match golden.",
                  fill=FILL_GREEN, align=ALIGN_CTR)
        else:
            for idx, exc in enumerate(all_exceptions, start=1):
                if exc.get("actual") == "missing" or exc.get("note", "").startswith("loan absent"):
                    row_fill = FILL_RED
                elif exc.get("actual") == "present" and exc.get("expected") == "absent":
                    row_fill = FILL_YELLOW
                else:
                    row_fill = FILL_ORANGE

                _cell(ws, current_row, 1, idx,                    fill=row_fill, align=ALIGN_CTR,  border=BORDER)
                _cell(ws, current_row, 2, exc.get("file", ""),    fill=row_fill, align=ALIGN_LEFT, border=BORDER)
                _cell(ws, current_row, 3, exc.get("loan_num",""), fill=row_fill, align=ALIGN_CTR,  border=BORDER)
                _cell(ws, current_row, 4, exc.get("column", ""),  fill=row_fill, align=ALIGN_LEFT, border=BORDER)
                _cell(ws, current_row, 5, exc.get("expected",""), fill=row_fill, align=ALIGN_LEFT, border=BORDER)
                _cell(ws, current_row, 6, exc.get("actual", ""),  fill=row_fill, align=ALIGN_LEFT, border=BORDER)
                _cell(ws, current_row, 7, exc.get("note",  ""),   fill=row_fill, align=ALIGN_LEFT, border=BORDER)
                current_row += 1

    wb.save(str(report_path))
    print(f"\nExcel report written: {report_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Funding & CashFlow regression harness: runs Final Funding (SG + CIBC) and "
            "CashFlow phases against each buy-date folder in TestData and diffs outputs "
            "against expected values.  Assumes _sg/_cibc exhibit files are already in "
            "files_required/ (tagging step already done)."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=r"""
Phases run per test case:
  1. Funding SG   (final_funding_sg.py)
  2. Funding CIBC (final_funding_cibc.py)
  3. CashFlow     (cashflow.compute.run_purchase_package for SG and CIBC)

Examples:
  python backend/scripts/regression_test_funding.py
  python backend/scripts/regression_test_funding.py --test-data C:\Users\omack\Downloads\TestData
  python backend/scripts/regression_test_funding.py --pdate 2026-02-24 --tday 2026-02-19
  python backend/scripts/regression_test_funding.py --no-cleanup
  python backend/scripts/regression_test_funding.py --report C:\temp\funding_results.xlsx
""",
    )
    parser.add_argument(
        "--test-data",
        type=str,
        default=r"C:\Users\omack\Downloads\TestData",
        help="Root directory containing buy-date test case folders.",
    )
    parser.add_argument(
        "--backend-dir",
        type=str,
        default=None,
        help="Path to the backend/ directory. Default: auto-detected.",
    )
    parser.add_argument(
        "--pdate",
        type=str,
        default=None,
        help="Purchase date YYYY-MM-DD. Applied to ALL test cases.",
    )
    parser.add_argument(
        "--tday",
        type=str,
        default=None,
        help="Base date YYYY-MM-DD for file naming. Applied to ALL test cases.",
    )
    parser.add_argument(
        "--no-cleanup",
        action="store_true",
        help="Do not delete temp work directories after each test case.",
    )
    parser.add_argument(
        "--update-golden",
        action="store_true",
        help=(
            "After running each test case, overwrite the golden expected files in "
            "TestData with the actual outputs.  Use this to re-baseline the golden data "
            "after a code change that intentionally changes outputs."
        ),
    )
    parser.add_argument(
        "--report",
        type=str,
        default=None,
        help="Path for the Excel report output. Default: FundingTestMatrix_<timestamp>.xlsx "
             "written to the test-data directory.",
    )
    args = parser.parse_args()

    test_data_dir  = Path(args.test_data)
    backend_dir    = Path(args.backend_dir) if args.backend_dir else _default_backend_dir()
    update_golden  = args.update_golden

    if not backend_dir.exists():
        print(f"ERROR: backend directory not found: {backend_dir}", file=sys.stderr)
        sys.exit(1)

    now_str   = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    print("=" * 60)
    print(f"FUNDING & CASHFLOW REGRESSION TEST — {now_str}")
    print(f"Test data: {test_data_dir}")
    if update_golden:
        print("MODE: --update-golden  (overwriting golden expected files with actual outputs)")
    print("=" * 60)

    test_cases   = discover_test_cases(test_data_dir)
    dates_config = _load_dates_config(test_data_dir)

    if not test_cases:
        print("\nNo valid test cases found. Nothing to run.")
        print("\n" + "=" * 60)
        print("SUMMARY: 0 PASSED / 0 FAILED / 0 TOTAL")
        print("=" * 60)
        sys.exit(0)

    results = []
    for case_dir in test_cases:
        result = run_test_case(
            test_case_dir=case_dir,
            backend_dir=backend_dir,
            cli_pdate=args.pdate,
            cli_tday=args.tday,
            no_cleanup=args.no_cleanup,
            dates_config=dates_config,
            update_golden=update_golden,
        )
        results.append(result)

    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] != "PASS")
    total  = len(results)

    print("\n" + "=" * 60)
    print(f"SUMMARY: {passed} PASSED / {failed} FAILED / {total} TOTAL")
    print("=" * 60)

    if failed > 0:
        print("\nFailed cases:")
        for r in results:
            if r["status"] != "PASS":
                err_note    = f" — {r['error']}" if r["error"] else ""
                phases_note = (
                    f" [failed phases: {', '.join(r['phases_failed'])}]"
                    if r["phases_failed"] else ""
                )
                issues_note = (
                    f" [{len(r['diffs'])} differ, {len(r['missing'])} missing, "
                    f"{len(r['extra'])} extra]"
                )
                print(f"  FAILED: {r['name']}{err_note}{phases_note}{issues_note}")

    # Write Excel report
    if args.report:
        report_path = Path(args.report)
    else:
        report_path = test_data_dir / f"FundingTestMatrix_{timestamp}.xlsx"

    try:
        write_excel_report(results, report_path)
    except ImportError:
        print("\n[WARN] openpyxl not installed — skipping Excel report. "
              "Install with: pip install openpyxl")

    sys.exit(0 if failed == 0 else 1)


if __name__ == "__main__":
    main()
